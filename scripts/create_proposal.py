from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports' / 'phase1_project_proposal.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
EV = json.loads((ROOT / 'outputs' / 'evaluation_summary.json').read_text())
MEAS = json.loads((ROOT / 'outputs' / 'measurement_summary.json').read_text())

NAVY = '173B56'
BLUE = '2F6685'
PALE = 'EAF2F7'
PALE2 = 'F5F8FA'
INK = '263746'
MUTED = '5A6E7C'


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, color=INK, size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, True, 'FFFFFF', size)
        shade(t.rows[0].cells[j], NAVY)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell(cells[j], val, False, INK, size)
            if ri % 2 == 1:
                shade(cells[j], PALE2)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.4)
styles['Normal'].font.color.rgb = RGBColor.from_string(INK)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for sty, sz, col in [('Title', 29, NAVY), ('Heading 1', 17, NAVY), ('Heading 2', 12.5, BLUE), ('Heading 3', 10.5, BLUE)]:
    s = styles[sty]
    s.font.name = 'Aptos Display' if sty != 'Heading 3' else 'Aptos'
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(col)
    s.paragraph_format.space_before = Pt(9)
    s.paragraph_format.space_after = Pt(4)

header = sec.header.paragraphs[0]
header.text = 'ENERGYMIND TUNISIA  |  PHASE 1 PROJECT PROPOSAL'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.name = 'Aptos'
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
footer = sec.footer.paragraphs[0]
footer.text = 'WattWitness · Microgrid Data Trust Layer · Phase 1'
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in footer.runs:
    r.font.name = 'Aptos'
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(62)
r = p.add_run('ENERGYMIND TUNISIA')
r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(24)
p.add_run('WattWitness')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Microgrid Data Trust Layer'); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(12)
r = p.add_run('Formal Project Proposal · Phase 1: Ideation & Data Exploration'); r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(MUTED)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(20)
p.add_run('Chosen system slice: one simulated microgrid node\nPV + wind + variable synthetic building load + battery + grid connection').bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(44)
p.add_run('Version 1.0  |  15 October 2026')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(40)
r = p.add_run('Proposal principle'); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph('Build and demonstrate a transparent, human-supervised trust layer between simulated energy telemetry and any later automated decision. Report what is implemented, what remains a design, and what still requires validation.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Executive summary and scope
doc.add_page_break()
doc.add_heading('1. Executive summary', level=1)
doc.add_paragraph('WattWitness is a Phase 1 prototype for the EnergyMind Tunisia challenge. It treats every telemetry reading as testimony and cross-examines it against the physical relationships it must satisfy and against simple data-quality rules. Each alert carries a plain-language reason and an evidence score and goes to a human reviewer, whose decision is recorded in an audit log. The goal is not to replace an operator: the system should state what its evidence supports and escalate uncertain cases.')
doc.add_paragraph('The prototype uses a one-year hourly simulation-derived dataset. PV and wind values are retained from a supplied base simulation export; the building load and context labels are synthetic; grid/battery dispatch is recalculated under stated assumptions; and voltage/current/inverter channels are derived through simplified formulas. A meter model then adds realistic measurement error, finite resolution and telemetry outages, so the data are not artificially perfect. Controlled faults are injected into a held-out window to test detection. Those results are simulated test results, not evidence of real-world performance.')
doc.add_heading('2. Scope', level=1)
add_table(doc, ['In scope for Phase 1', 'Out of scope / future work'], [
    ('One simulated microgrid telemetry node; hourly series for PV, wind, building load, grid import, battery and derived electrical signals.', 'Production deployment, real-time connection to utility infrastructure, autonomous equipment control, or automatic data repair.'),
    ('Controlled injection of sensor/data-quality faults with traceable event metadata.', 'Claims that injected events replicate every real equipment failure or cyberattack.'),
    ('Transparent rules and physical-consistency checks; comparisons against baselines.', 'A fully validated ML model or calibrated probability of fault.'),
    ('Human review queue with recorded decisions and an escalation policy.', 'Multi-node exchange or federated learning implementation; Phase 1 includes design documentation only.'),
], widths=[3.25, 3.25], size=8.2)
doc.add_heading('3. Objectives', level=1)
for x in [
    'Detect selected classes of missing, abrupt, frozen, delayed, drifting and biased telemetry in a repeatable experiment.',
    'Compare suspicious readings with related power, voltage and current signals to surface inconsistencies.',
    'Give each alert a plain-language rationale, evidence strength, recommended check and human-review state.',
    'Define a human review step in which decisions and reasons are recorded and no equipment is controlled.',
    'Define clear escalation boundaries and propose a privacy-preserving mechanism for future sharing between nodes.',
    'Deliver the prepared dataset, source code, experiment results, technical report and documented limitations.'
]: bullet(doc, x)

# Dataset plan
doc.add_page_break()
doc.add_heading('4. Data and anomaly-simulation plan', level=1)
doc.add_heading('4.1 Baseline dataset', level=2)
doc.add_paragraph('The prepared dataset contains 8,784 hourly rows from 1 June 2023 through 31 May 2024. The renewable PV/wind series are retained from the supplied simulation export. The original repeated daily load is used only as a rough shape template; a different synthetic building load is generated with weekday/weekend differences, seasonal multipliers, daily variability, correlated hourly noise and intermittent demand events. EV load is excluded.')
doc.add_paragraph(f'Because derived channels satisfy their equations exactly, a seeded meter model converts this ideal state into measured telemetry: about 0.5% error on power and current, 0.2% on voltage, finite reporting resolution, and {MEAS["natural_gap_events"]} natural outages leaving {MEAS["rows_with_any_missing_value"]} rows with unknown values. Detection limits are learned from the training period of this measured data.')
add_table(doc, ['Signal family', 'Construction and interpretation'], [
    ('PV and wind power', 'Retained from the supplied base simulation series; units and measurement/simulation provenance are inherited and still need confirmation.'),
    ('Building load and context', 'Synthetic variable building demand plus generated context labels. Not measured demand or activity.'),
    ('Grid and battery', 'Recalculated using a stated balance equation, 100 kW battery limit, assumed 100–500 kWh stored-energy range and 95% charge/discharge efficiency.'),
    ('Voltage, current and inverter', 'Derived with simplified formulae using assumed 400 V three-phase voltage and power factors. These are not real sensor measurements or detailed equipment simulations.'),
], widths=[1.55, 4.95], size=8.2)
doc.add_heading('4.2 Injection protocol', level=2)
doc.add_paragraph(f'The final 60 days of the prepared dataset are held out. A fixed seed creates 18 traceable events (three per type, from gross to subtle severity) altering {EV["labelled_fault_rows"]} hourly readings; only readings whose value actually changed are labelled. Each injected event records its type, target signal, start/end times, duration, severity and seed. The untouched baseline is retained separately; injection labels are excluded from the detector inputs and are used only to evaluate predictions.')
add_table(doc, ['Injected scenario', 'Demonstration intent'], [
    ('Spike in grid power', 'Detect a one-sample extreme jump.'),
    ('Frozen PV reading', 'Detect a meaningful signal remaining unchanged for too long.'),
    ('Gradual building-load drift', 'Test a slow offset against the modelled cross-signal relationships.'),
    ('Missing battery SOC', 'Detect missing telemetry values.'),
    ('Delayed/replayed PV signal', 'Test a reading that is plausible in magnitude but misaligned in time.'),
    ('Plausible grid-current bias', 'Test an altered measurement that can conflict with voltage/power consistency.'),
], widths=[2.2, 4.3], size=8.2)
doc.add_heading('4.3 Evaluation', level=2)
doc.add_paragraph('Compare simple rules, physical-consistency checks, their transparent combination, and an Isolation Forest baseline on the same held-out test window. Report precision, recall, F1 and average precision, event-level detection, recall by fault type, alerts on the untouched window, and sensitivity to meter error. The untouched-window figure is described as an alert rate, not a verified false-positive rate, because there are no independent labels for every real event in the base simulation.')

# four goals exact
doc.add_page_break()
doc.add_heading('5. Preliminary design for each official Challenge Goal', level=1)
doc.add_paragraph('The four goal names below follow the official EnergyMind Tunisia challenge description. The implementation status is stated conservatively; a preliminary design is not the same as a completed implementation.')
doc.add_heading('Goal 1 — Anomaly & Bias Detection', level=2)
doc.add_paragraph('Design: monitor the available energy variables and derive checks for missing values, abrupt jumps, frozen values and disagreement between linked electrical channels. The hybrid alert is raised when a data-quality rule triggers or a physical-consistency residual exceeds the current threshold. The controlled experiment covers spike, frozen value, gradual drift, missing value, delayed/replayed signal and plausible current bias.')
doc.add_paragraph('Distinguishing event types: a reading that conflicts with related channels supports a data-quality hypothesis. A coherent change across multiple channels without a strong physical mismatch is shown as a potential operating event requiring context confirmation. In Phase 1 the system cannot establish causality with certainty because the simulation lacks independent labels for genuine grid events and the synthetic context is not a trusted operational log. Power-quality waveforms and sub-second smart-meter streams are outside the included hourly dataset, so the present implementation focuses on the signals actually available.')
doc.add_paragraph('Status: partially implemented. The single-node detector and controlled test are implemented. Context-aware discrimination between genuine events, natural equipment failures and telemetry corruption requires stronger context data and independently labelled cases.')

doc.add_heading('Goal 2 — Confidence Scores & Explainability', level=2)
doc.add_paragraph('Design: each alert carries a plain-language reason, the related residuals/rules, a qualitative evidence label, a suggested operator check, and an audit-row timestamp. Each alert also carries a simple heuristic evidence score and the size of every mismatch relative to its learned limit. The score is explicitly labelled as uncalibrated and is not presented as the probability that a fault occurred.')
doc.add_paragraph('Status: partially implemented. Explanations and logging are generated. Confidence calibration, SHAP/LIME analysis, and independent human scoring of explanation quality are not yet completed.')

doc.add_heading('Goal 3 — Human-AI Escalation & Governance', level=2)
doc.add_paragraph('Design: the AI may flag and explain, but will not change raw readings or control equipment. A human reviewer can confirm a data issue, judge that a change may be a genuine operating event, dismiss an alert, or escalate it, while recording the reason. Decisions are appended to an audit log whose existing entries are never rewritten. To reduce over-trust, the score is labelled as heuristic and shows the evidence and limitations; to reduce under-trust, it makes the checks and supporting signals inspectable and gives the operator a documented override/escalation path.')
add_table(doc, ['Provisional evidence tier', 'Proposed response in Phase 1'], [
    ('No meaningful rule or physical conflict; heuristic score below 0.35', 'No automatic action. Continue monitoring; do not treat the lack of an alert as proof of safety.'),
    ('Rule pattern without a strong physical mismatch; score around 0.35–0.69', 'Queue for operator review. Ask the operator to inspect adjacent values, timing and context.'),
    ('Strong physical mismatch; score 0.70 or higher', 'Prioritize human review and compare related signals/source records. Do not auto-correct or trip equipment.'),
    ('Missing critical signal, contradictory evidence or uncertain diagnosis', 'Escalate to an authorized operator; preserve the raw data and log the decision.'),
], widths=[2.5, 4.0], size=8.0)
doc.add_paragraph('These score bands are illustrative governance thresholds for the demo, not calibrated values or a safety-certified policy. Automatic corrective actions remain disabled. Status: review workflow prototyped; production governance remains future work.')

doc.add_heading('Goal 4 — Shared Learning Across the Noosphere', level=2)
doc.add_paragraph('Design: nodes may share approved anonymized anomaly patterns, aggregated calibration insights and relevant knowledge as normalized summaries (signal family, coarse shape/duration, checks triggered, detector version and reviewed outcome), not raw traces, precise timestamps, device identifiers or private schedules by default. Each receiving node checks units, signals, sign conventions and physical compatibility; patterns run in shadow mode and require provenance and review before adoption. The design considers privacy leakage, poisoned knowledge, negative transfer and false confidence.')
doc.add_paragraph('Status: design documented only. No multi-node exchange, shared model or federated training is implemented in Phase 1; the challenge states that multi-node sharing may remain a documented design.')

# Architecture, implementation matrix
doc.add_heading('6. System architecture and operator workflow', level=1)
for item in [
    'Prepare the one-year simulated dataset; leave the clean source table unchanged.',
    'Apply the meter model to obtain measured telemetry with realistic error and gaps.',
    'Create a labeled faulted copy of the held-out test window using fixed injection settings.',
    'Run data-quality rules and simplified physical-consistency checks; compare against baselines.',
    'Generate alert explanations, evidence labels, metrics and a review queue.',
    'A reviewer inspects each alert with its related signals and records a decision and a reason.',
    'Retain a local audit record of reviewer label, decision and notes. No equipment command is sent.'
]: numbered(doc, item)
doc.add_heading('7. Implementation status and deliverables', level=1)
add_table(doc, ['Phase 1 element', 'Status'], [
    ('Prepared simulated dataset and data dictionary', 'Included; base export is not included in the public bundle. Confirm rights for derived PV/wind series before public upload.'),
    ('Fault injection, detector and evaluation scripts', 'Implemented; repeatable with included prepared dataset.'),
    ('Metric output and plots', 'Included; results are on controlled synthetic faults only.'),
    ('Explainable alert output and review queue', 'Implemented; evidence score is heuristic, not calibrated.'),
    ('Human review queue and decision log', 'Implemented as a prototype; not a secured production audit system.'),
    ('Cross-node privacy/sharing approach', 'Documented design only, as permitted for Phase 1.'),
    ('Formal proposal and technical report', 'Included as Word and PDF.'),
], widths=[2.55, 3.95], size=8.0)

doc.add_heading('8. Main risks and next validation steps', level=1)
for x in [
    'Verify inherited PV/wind provenance, units and redistribution rights for the prepared dataset.',
    'Verify the battery-state interpretation, dispatch assumptions and sign convention against the originating simulation model.',
    'Add examples of coherent genuine operating events and independently labelled sensor faults; evaluate event-versus-data classification.',
    'Calibrate evidence scores on labelled data and assess explanation quality with operators.',
    'If authorized, test replayed or non-invasive data with an independent laboratory node; do not connect the prototype to live control equipment.',
    'For multi-node learning, implement sender identity, versioned schema, provenance checks, privacy tests and shadow-mode evaluation before activation.'
]: bullet(doc, x)
# Core properties without personal metadata
doc.core_properties.title = 'WattWitness — Phase 1 Project Proposal'
doc.core_properties.subject = 'Simulation-based microgrid data-trust prototype and four-goal design'
doc.core_properties.author = 'WattWitness team'
doc.core_properties.keywords = 'EnergyMind, simulated microgrid, anomaly detection, explainability, human review, privacy'
doc.save(OUT)
print(OUT)
