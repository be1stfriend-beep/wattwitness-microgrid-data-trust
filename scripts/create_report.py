#!/usr/bin/env python3
"""Build the Phase 1 technical report (reports/phase1_report.docx) from the pipeline outputs."""
from pathlib import Path
from collections import Counter
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTS = ROOT / 'outputs'
REPORT = ROOT / 'reports' / 'phase1_report.docx'
REPORT.parent.mkdir(parents=True, exist_ok=True)
INK, AMBER, TEAL, MUTED, PALE = '10212E', 'C27C0E', '0F766E', '5B6B78', 'F3F1EC'
DATE = '15 October 2026'

audit = json.loads((OUTS / 'dataset_audit_summary.json').read_text())
meas = json.loads((OUTS / 'measurement_summary.json').read_text())
ev = json.loads((OUTS / 'evaluation_summary.json').read_text())
metrics = pd.read_csv(OUTS / 'evaluation_metrics.csv')
by_type = pd.read_csv(OUTS / 'fault_type_recall.csv')
events = pd.read_csv(OUTS / 'event_level_results.csv')
rates = pd.read_csv(OUTS / 'untouched_test_alarm_rates.csv')
noise = pd.read_csv(OUTS / 'noise_robustness.csv')
manifest = pd.read_csv(OUTS / 'fault_manifest.csv')

sel = metrics[metrics.detector.str.contains('selected')].iloc[0]
iso = metrics[metrics.detector.eq('Isolation Forest')].iloc[0]
full = metrics[metrics.detector.str.startswith('Full hybrid')].iloc[0]
rules_only = metrics[metrics.detector.eq('Rules only')].iloc[0]
phys_only = metrics[metrics.detector.eq('Physical checks only')].iloc[0]
rate = rates[rates.detector.str.contains('selected')].iloc[0]
n_events, n_detected, n_full = len(events), int(events.detected.sum()), int(events.fully_detected.sum())
n_fault_rows = int(ev['labelled_fault_rows'])
nz = lambda scale, kind: noise[(noise.noise_scale == scale) & noise.thresholds.str.startswith(kind)].iloc[0]
fixed1, cal1, cal4 = nz(1.0, 'fixed'), nz(1.0, 'limits'), nz(4.0, 'limits')
gap_devices = Counter(g['device'] for g in meas['gaps'])
TH = ev['physics_thresholds']


# ------------------------------------------------------------------ helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def cell_text(cell, text, bold=False, color=INK, size=8.8):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(par, instr, placeholder=''):
    for kind in ('begin', 'instr', 'separate', 'text', 'end'):
        r = par.add_run()
        if kind == 'text':
            r.text = placeholder
            continue
        if kind == 'instr':
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = instr
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind)
        r._r.append(el)


counter = {'Table': 0, 'Figure': 0}


def caption(kind, text):
    counter[kind] += 1
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == 'Figure' else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f'{kind} {counter[kind]}. '); r.bold = True
    p.add_run(text)


def table(headers, rows, widths, title, size=8.8):
    caption('Table', title)
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell_text(t.rows[0].cells[j], h, True, 'FFFFFF', size); shade(t.rows[0].cells[j], INK)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cell_text(cells[j], v, False, INK, size)
            if i % 2:
                shade(cells[j], PALE)
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, title, width=6.3):
    doc.add_picture(str(ROOT / 'figures' / name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption('Figure', title)


def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(text); p.alignment = align
    return p


def bullet(text, lead=None):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    if lead:
        p.add_run(lead + ' ').bold = True
    p.add_run(text)


def equation(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = True; r.font.size = Pt(10)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def pct(x):
    return f'{100 * x:.1f}%'


# ------------------------------------------------------------------ document setup
doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.85), Inches(.8)
sec.left_margin = sec.right_margin = Inches(.95)
sec.different_first_page_header_footer = True
st = doc.styles
st['Normal'].font.name = 'Calibri'; st['Normal'].font.size = Pt(10.5)
st['Normal'].font.color.rgb = RGBColor.from_string('23303A')
st['Normal'].paragraph_format.space_after = Pt(6); st['Normal'].paragraph_format.line_spacing = 1.12
for name, size, color in [('Title', 30, INK), ('Heading 1', 16, INK), ('Heading 2', 12.5, TEAL), ('Heading 3', 11, TEAL)]:
    s = st[name]; s.font.name = 'Cambria'; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 9); s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True
st['Caption'].font.size = Pt(9); st['Caption'].font.italic = False; st['Caption'].font.color.rgb = RGBColor.from_string(MUTED)
hp = sec.header.paragraphs[0]; hp.text = 'WattWitness · Phase 1 technical report'; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(MUTED)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('EnergyMind Tunisia · Phase 1 · page '); fr.font.size = Pt(8.5); fr.font.color.rgb = RGBColor.from_string(MUTED)
add_field(fp, 'PAGE', '1')

# ------------------------------------------------------------------ cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(70)
r = p.add_run('ENERGYMIND TUNISIA · PHASE 1 — IDEATION & DATA EXPLORATION'); r.bold = True; r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(AMBER)
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(26)
p.add_run('WattWitness')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Cross-examining microgrid telemetry:\nexplainable detection of untrustworthy energy data'); r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
r = p.add_run('Technical report · Version 1.0 · ' + DATE); r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(MUTED)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(120)
r = p.add_run('Scope of this report'); r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
para('One simulated microgrid node (PV, wind, synthetic building load, battery and grid connection), hourly data for one year. '
     'All results come from controlled fault injection on simulated telemetry with modelled meter error; they validate the method, '
     'not performance on physical equipment.', WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ------------------------------------------------------------------ abstract + contents
h1('Abstract')
para(f'Energy-management decisions are only as good as the telemetry they rely on, yet meters drift, freeze, drop out, '
     f'and can be tampered with. WattWitness is a data-trust layer that treats every reading as testimony and cross-examines it '
     f'against the physical relationships it must satisfy (power balance, voltage × current × power factor, battery current) '
     f'and against simple data-quality rules. A one-year hourly simulation of a single microgrid node was built from a supplied '
     f'PV/wind series, a variable synthetic building load and recomputed dispatch; a meter model then added realistic measurement '
     f'error, finite resolution and {meas["natural_gap_events"]} natural telemetry outages so that the data are not artificially perfect. '
     f'Detection limits were learned from the first {ev["training_rows"]:,} hours only. On a held-out 60-day window containing '
     f'{n_events} injected faults of six types ({n_fault_rows} altered hourly readings), the selected rules + physics detector reached '
     f'precision {sel.precision:.3f}, recall {sel.recall:.3f} and F1 {sel.f1:.3f}; it flagged {n_detected} of {n_events} events and raised '
     f'alerts on {rate.alarm_rate_pct:.2f}% of unmodified hours. An Isolation Forest baseline reached F1 {iso.f1:.3f}. A sensitivity '
     f'study shows that fixed thresholds collapse once meter error is present (F1 {fixed1.f1:.2f} at 0.5% error) while learned limits '
     f'degrade gracefully (F1 {cal4.f1:.2f} at 2% error). Each alert carries the checks that fired, the size of every mismatch '
     f'relative to its limit, a transparent evidence score and a suggested check, and is routed to a human reviewer; nothing is '
     f'corrected automatically.')
p = doc.add_paragraph(); r = p.add_run('Keywords: '); r.bold = True
p.add_run('microgrid telemetry, data quality, physics-based consistency checks, anomaly detection, explainability, human oversight.')
h1('Contents')
add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', 'Right-click and choose "Update field" to build the table of contents.')
doc.add_page_break()

# ------------------------------------------------------------------ 1 introduction
h1('1. Introduction')
h2('1.1 Context and motivation')
para('Microgrids combine variable renewable generation, storage and a grid connection, and are increasingly operated by software '
     'that reads telemetry and acts on it. A wrong reading — a stuck PV meter, a biased current transformer, a delayed or replayed '
     'value, a communication outage — can silently propagate into dispatch, billing and maintenance decisions. Because a genuine '
     'operating change can look just as unusual as a faulty sensor, simply flagging outliers is not enough: an operator needs to know '
     'which readings disagree with each other and why.')
h2('1.2 Problem statement')
para('Given hourly telemetry from one microgrid node, identify readings that are unlikely to be trustworthy, explain the evidence '
     'in plain language, quantify how strong that evidence is, and hand the decision to a human — without ever correcting data '
     'or controlling equipment automatically.')
h2('1.3 Objectives and scope of Phase 1')
for t in ['Prepare and explore a reproducible one-year dataset for one simulated node, with honest provenance for every signal.',
          'Make the telemetry realistically imperfect (meter error, resolution, outages) so that checks are not validated on perfect data.',
          'Detect six fault classes — spike, frozen value, gradual drift, missing value, delayed/replayed signal and plausible bias — '
          'with transparent rules and physics cross-checks, and compare them with a statistical baseline.',
          'Attach an explanation, an evidence score and a review status to every alert, and define the escalation policy.',
          'Document a privacy-preserving design for sharing knowledge between nodes (Goal 4), as permitted for Phase 1.']:
    bullet(t)
h2('1.4 Contributions')
for lead, t in [('Measured-telemetry layer.', 'A seeded meter model that separates the ideal simulated state from what meters would report.'),
                ('Learned physics limits.', 'Six residual checks whose limits are learned from the training period, not hand-picked.'),
                ('Honest evaluation.', 'Faults of decreasing severity, labels only on readings that really changed, event-level and '
                 'row-level metrics, a clean-window alert rate and a noise-sensitivity study.'),
                ('Explainable, human-governed alerts.', 'Every alert states which checks fired and by how much, and requires a human decision.')]:
    bullet(t, lead)

# ------------------------------------------------------------------ 2 background
h1('2. Background')
para('Residual-based bad-data detection is a long-standing idea in power-system state estimation: measurements that are '
     'inconsistent with the network model produce large residuals and are investigated [1]. WattWitness applies the same principle at '
     'the scale of one node, using the power balance and the electrical relationships between power, voltage and current. '
     'Unsupervised statistical detectors such as Isolation Forest [2] need no physical model and are used here as a comparison '
     'baseline [3]. Realistic meter error is bounded by accuracy classes such as IEC 62053-22 class 0.5S [4], which motivates the '
     'meter model in Section 3.4. Finally, risk-management guidance for AI systems stresses transparency and meaningful human '
     'oversight [5], which shapes the explanation and escalation design in Sections 4.5–4.6.')

# ------------------------------------------------------------------ 3 data
h1('3. Data')
h2('3.1 Source and selected window')
para(f'The base input is a supplied simulation export (believed to come from a PyPSA model [6]); it is not redistributed. The latest '
     f'complete year was selected: {audit["period_start"][:10]} to {audit["period_end"][:10]}, {audit["rows"]:,} hourly rows '
     f'(2024 is a leap year). The window has no duplicate timestamps and no missing hours. PV and wind series are kept as supplied; '
     f'their units are assumed to be kW and their provenance is inherited.')
h2('3.2 Synthetic building load and context')
para(f'The supplied load repeats one daily profile and cannot represent a real building, so it is used only as an hourly shape '
     f'template. The synthetic load varies by weekday/weekend, month, a random factor per day, correlated hourly noise and occasional '
     f'high-demand activities ({audit["scheduled_activity_hours"]} hours), ranging from {audit["load_min"]:.0f} to '
     f'{audit["load_max"]:.0f} kW with {audit["load_unique_rounded_2dp"]:,} distinct values. Context labels (weekday, season, activity) '
     f'are generated with it and are not measurements. EV load is excluded.')
figure('building_load_variability_14days.png', 'First 14 days: original repeated template vs the variable synthetic building load.')
h2('3.3 Dispatch and derived electrical signals')
para('Battery and grid flows are recomputed hour by hour for the new load; voltage, current and inverter channels are then derived '
     'from the powers with simplified equations. The assumptions are listed in Table 1 and must be confirmed before any physical claim.')
table(['Signal group', 'How it is produced', 'Status'], [
    ['PV and wind power', 'Copied from the selected dates of the supplied export', 'Source series; provenance inherited'],
    ['Building load and context', 'Template shape × weekday, seasonal, daily and hourly variation; generated labels', 'Synthetic'],
    ['Battery power and SOC', '100–500 kWh energy range (assumed), 100 kW limit, 95% efficiency, charge positive', 'Simulated'],
    ['Grid import and curtailment', 'Import only; renewable surplus the battery cannot absorb is curtailed', 'Simulated'],
    ['Inverter, voltage, current', '97% inverter efficiency, 400 V three-phase, PF 0.99 (PV) and 0.98 (grid), linear 350–420 V battery', 'Derived'],
], [1.55, 3.6, 1.45], 'Construction of each signal family and its status.')
h2('3.4 Measurement realism: meter error, resolution and telemetry gaps')
para('Derived from the same equations that are later checked, the ideal state satisfies every relationship exactly: the largest '
     f'power-balance residual over the year is {audit["power_balance_residual_abs_max"]:.1e} kW. Testing physics checks on such data '
     'would be circular. A seeded meter model therefore converts the ideal state into measured telemetry (Table 2): each channel '
     'receives Gaussian error proportional to its value plus a small floor, is rounded to its reporting resolution, and devices that '
     'report together drop out together. A device that is off still reports exactly zero.')
table(['Channels', 'Relative error (1σ)', 'Floor (1σ)', 'Resolution'], [
    ['Power channels (PV, wind, load, grid, battery, inverter AC/DC)', '0.5%', '0.1 kW', '0.1 kW'],
    ['Curtailment (estimated, not metered)', '1.0%', '0.1 kW', '0.1 kW'],
    ['Battery state of charge', '0.3%', '0.5 kWh', '0.1 kWh'],
    ['Voltages', '0.2%', '0.2 V', '0.1 V'],
    ['Currents', '0.5%', '0.2 A', '0.1 A'],
], [3.2, 1.2, 1.0, 1.0], 'Meter model (seed 7203). Relative errors follow class 0.5 metering practice [4].')
para(f'Outages start with probability {meas["dropout_start_probability_per_hour"]:.3f} per hour and last one to three hours. The year '
     f'contains {meas["natural_gap_events"]} outages ({", ".join(f"{v} {k}" for k, v in gap_devices.most_common())}), leaving '
     f'{meas["rows_with_any_missing_value"]} rows with at least one unknown value ({meas["missing_cells"]} empty cells). Figure 2 '
     f'shows the effect on the power-balance residual: the ideal state sits exactly on zero, the measured telemetry does not.')
figure('power_balance_residual.png', 'Power-balance residual over the year: ideal state vs measured telemetry, with the learned limit.')
h2('3.5 Exploration summary')
figure('annual_power_profiles_14day_mean.png', 'Annual profiles (14-day rolling mean) of the main power flows.')
para('PV dominates in spring and summer, the synthetic load peaks in the summer cooling season, and grid import fills the winter '
     'deficit. These seasonal shifts matter for the detector: a large hourly change is often a legitimate dispatch step, which is '
     'why the spike rule in Section 4.2 requires a jump that reverses rather than any large jump.')

# ------------------------------------------------------------------ 4 method
h1('4. Method')
h2('4.1 Architecture')
figure('architecture.png', 'Single-node processing chain. The Isolation Forest is evaluated for comparison only.', 6.5)
para('The ideal state is turned into measured telemetry; limits are learned on the training year; the test window (with injected '
     'faults) is screened by the rules and the physics checks; each alert is explained and queued for a human decision.')
h2('4.2 Data-quality rules')
bullet('any of the 17 telemetry channels has no value.', 'Missing:')
bullet('a channel jumps by more than a fixed step (e.g. 160 kW for grid power) and the next sample moves back by at least half of '
       'that step. A step that persists — typical of battery and grid dispatch — is not flagged. The rule confirms a spike one hour late.',
       'Spike-and-return:')
bullet('a non-zero channel repeats exactly the same value for four consecutive hours. With real meter error, exact repeats of a '
       'non-zero value are very unlikely unless the sensor is stuck.', 'Frozen:')
h2('4.3 Physics cross-checks')
para('With battery power positive while charging and grid power non-negative for imports, six residuals are computed per hour:')
equation('r₁ = P_PV,AC + P_wind + P_grid − P_load − P_battery − P_curtailed')
equation('r₂ = P_PV,DC − V_DC·I_DC      r₃ = P_PV,AC − √3·V_LL·I·PF_PV      r₄ = P_grid − √3·V_LL·I·PF_grid')
equation('r₅ = I_battery − I_expected(P_battery, V_battery, η)      r₆ = P_PV,meter − P_PV,inverter')
para('A check fires when |r| exceeds its limit. Each limit is the 99.9th percentile of |r| over the training year plus a 20% margin '
     '(Table 3), so it reflects the meter error actually present rather than a hand-picked number.')
names = {'power_balance_kw': 'r₁ power balance', 'pv_dc_kw': 'r₂ PV DC power', 'pv_ac_kw': 'r₃ PV AC power', 'grid_kw': 'r₄ grid power',
         'battery_current_a': 'r₅ battery current', 'pv_source_vs_inverter_kw': 'r₆ PV meter vs inverter'}
table(['Check', 'Learned limit', 'Fixed v0.1 limit'],
      [[names[k], f'{v:.2f} {"A" if k.endswith("_a") else "kW"}', '5 A' if k.endswith('_a') else '2 kW'] for k, v in TH.items()],
      [2.6, 1.6, 1.6], 'Physics limits learned on the training period.')
h2('4.4 Statistical baseline')
para('An Isolation Forest [2] (160 trees, median imputation, robust scaling) is fitted on the 17 channels of the training period. '
     'Its alert threshold is the 98th percentile of its anomaly score over the last 1,000 training hours. It is reported as a '
     'baseline and as part of a "full hybrid" variant, not used in the selected detector.')
h2('4.5 Explanation and evidence score')
para('Every alert lists the checks that fired with their values, e.g. "Grid power vs grid voltage × current off by 12.2 kW (limit 3.9 kW)", '
     'together with an interpretation and a suggested check. A transparent evidence score summarises the strength of the evidence '
     '(Table 4). It is a heuristic, not a calibrated probability.')
table(['Situation', 'Interpretation', 'Score', 'Policy'], [
    ['A channel has no value', 'Telemetry gap', '0.90', 'Check logger / link; never fill silently'],
    ['Physics residual above limit (ratio k = |r|/limit)', 'Likely telemetry / data inconsistency', '0.70 + 0.05(k−1), max 0.95', 'Prioritised human review'],
    ['Rule only, several power channels moved together', 'Potential operating event', '0.45', 'Verify context before calling it a fault'],
    ['Rule only, no coherent change', 'Possible data-quality issue', '0.60', 'Queue for review'],
    ['No check fired', 'No integrity conflict detected', '0.10', 'Monitor; absence of alert is not proof'],
], [2.0, 1.75, 1.25, 1.6], 'Interpretation, evidence score and escalation policy.')
h2('4.6 Human review and escalation')
para('Each alert enters a review queue with status PENDING_REVIEW. A reviewer records one of four decisions — confirm data issue, '
     'likely genuine operating event, escalate, or dismiss — with a reason. Decisions are appended to an audit log keyed by scenario '
     'and timestamp; existing entries are never rewritten. The system never modifies raw readings and never sends a command to equipment.')

# ------------------------------------------------------------------ 5 experimental design
h1('5. Experimental design')
h2('5.1 Training and test split')
para(f'The first {ev["training_rows"]:,} hours (about 84%) form the training period, used only to learn limits and fit the baseline. '
     f'The final {ev["test_rows"]:,} hours ({ev["test_period_start"][:10]} to {ev["test_period_end"][:10]}) are held out. A clean copy '
     f'of this window is kept; faults are injected into a second copy.')
h2('5.2 Fault-injection protocol')
para('Eighteen non-overlapping events (three per type) are scheduled with a fixed seed. Spikes and biases become smaller in each '
     'round so that the test probes the detection limit (Table 5). Injections never overlap a natural outage. Only readings whose '
     'value actually changed are labelled as faults: the first hour of a drift (zero offset) or of a frozen sensor (holding the true '
     'value) is indistinguishable from clean data and is not counted.')
sev = manifest.groupby('fault_type')['severity'].apply(lambda s: ' / '.join(dict.fromkeys(x.split(' of ')[-1] if 'spike' in x else x for x in s)))
spec = {'spike': ('Grid Power', '1 h'), 'frozen': ('PV Power', '8 h'), 'gradual_drift': ('Building Load', '24 h'),
        'missing': ('Battery SOC', '4 h'), 'delayed_signal': ('PV Power', '12 h'), 'plausible_bias': ('Grid Current', '6 h')}
table(['Fault type', 'Signal', 'Duration', 'Severity (round 1 / 2 / 3)', 'Altered rows'],
      [[k, s, d, sev.get(k, ''), int(manifest[manifest.fault_type == k].altered_rows.sum())] for k, (s, d) in spec.items()],
      [1.15, 1.1, 0.7, 2.75, 0.9], 'Injected fault classes.', size=8.4)
h2('5.3 Metrics')
for lead, t in [('Row level:', 'precision, recall, F1 and average precision over the test hours. The few hours of the test window '
                 'that contain a natural outage are excluded from these metrics and reported separately.'),
                ('Event level:', 'an event is detected if at least one altered reading is flagged, and fully detected if all are.'),
                ('Clean-window alert rate:', 'alerts on the unmodified copy. It is not a verified false-positive rate, because the '
                 'simulation has no independent labels for every genuine event.'),
                ('Sensitivity:', 'the full experiment is repeated for meter error from 0% to 2%, with fixed and with learned limits.')]:
    bullet(t, lead)

# ------------------------------------------------------------------ 6 results
h1('6. Results')
h2('6.1 Detector comparison')
table(['Detector', 'Precision', 'Recall', 'F1', 'Avg. precision', 'TP', 'FP', 'FN'],
      [[r.detector, f'{r.precision:.3f}', f'{r.recall:.3f}', f'{r.f1:.3f}', f'{r.average_precision:.3f}', r.tp, r.fp, r.fn]
       for r in metrics.itertuples()], [2.45, .7, .62, .55, .82, .42, .42, .42], 'Row-level results on the held-out window with injected faults.', size=8.3)
para(f'The selected rules + physics detector found {sel.tp} of {n_fault_rows} altered readings with {sel.fp} false alerts '
     f'(precision {sel.precision:.3f}, recall {sel.recall:.3f}, F1 {sel.f1:.3f}). Rules alone are precise but catch only '
     f'{pct(rules_only.recall)} of faulty readings; physics checks alone catch {pct(phys_only.recall)}; combining them adds missing '
     f'values and frozen readings that physics cannot see. The Isolation Forest is much weaker (F1 {iso.f1:.3f}) and adding it to '
     f'the selected detector lowers precision to {full.precision:.3f}, so it is not part of the selected configuration.')
figure('detector_comparison.png', 'Precision, recall and F1 per detector on injected faults (0.5% meter error).')
h2('6.2 Results by fault type and by event')
table(['Fault type', 'Altered rows', 'Rules recall', 'Physics recall', 'Rules + physics recall', 'Events detected'],
      [[r.fault_type, r.affected_rows, f'{r.rules_recall:.2f}', f'{r.physics_recall:.2f}', f'{r.hybrid_recall:.2f}', f'{r.events_detected} / {r.events}']
       for r in by_type.itertuples()], [1.35, .95, .9, .95, 1.4, 1.05], 'Recall by fault type.')
para(f'All {n_events} events were flagged at least once and {n_full} were flagged on every altered hour. The misses are '
     f'explainable: the first hours of each gradual drift stay inside meter error until the offset exceeds the '
     f'{TH["power_balance_kw"]:.1f} kW power-balance limit (first flag after {int(events[events.fault_type == "gradual_drift"].hours_to_first_flag.min())}–'
     f'{int(events[events.fault_type == "gradual_drift"].hours_to_first_flag.max())} hours), and the subtlest 5% current bias is '
     f'only visible in the hours with high grid current. Every gross and medium-severity spike and bias was caught.')
h2('6.3 Unmodified window')
para(f'On the clean copy the selected detector raised {int(rate.alarm_rows)} alerts in {int(rate.untouched_test_rows):,} hours '
     f'({rate.alarm_rate_pct:.2f}%); the {int(rate.natural_gap_rows)} hours with a natural outage were also flagged, correctly, as '
     f'telemetry gaps. The alert rate on the training period is {ev["training_alert_rate_pct"]:.2f}%, almost the same, which indicates '
     f'that the limits were not tuned to the test window.')
h2('6.4 Robustness to meter error')
table(['Meter error (1σ)', 'Limits', 'Precision', 'Recall', 'F1', 'Clean-hour alert rate'],
      [[f'{r.power_meter_error_pct:.2f}%', 'fixed (v0.1)' if r.thresholds.startswith('fixed') else 'learned', f'{r.precision:.3f}',
        f'{r.recall:.3f}', f'{r.f1:.3f}', f'{r.clean_window_alert_rate_pct:.2f}%'] for r in noise.itertuples()],
      [1.15, 1.0, .85, .75, .7, 1.45], 'Sensitivity of the rules + physics detector to meter error.', size=8.4)
para(f'Fixed 2 kW / 5 A limits only work on noise-free data: at the realistic 0.5% error they flag {fixed1.clean_window_alert_rate_pct:.1f}% '
     f'of clean hours and F1 falls to {fixed1.f1:.2f}. Learned limits keep the clean-hour alert rate around '
     f'{cal1.clean_window_alert_rate_pct:.1f}% at every noise level; the price is lower recall for subtle faults as error grows '
     f'(F1 {cal4.f1:.2f} at 2%). On noise-free data the frozen rule also fires on legitimate plateaus such as a battery charging at '
     f'its 100 kW limit — another artefact of perfect data that disappears with realistic meters.')
figure('noise_robustness.png', 'F1 and clean-hour alert rate as meter error grows, with fixed and learned limits.')
h2('6.5 Worked example')
figure('fault_case_grid_current_bias.png', 'A 12% grid-current bias: the reported current departs from the clean value and every biased hour is flagged.', 6.2)
para('The biased current is still plausible on its own, so no data-quality rule fires. The physics check exposes it: grid power '
     'no longer equals √3·V·I·PF, by about 12% of the grid power. The alert names the three grid channels to compare.')

# ------------------------------------------------------------------ 7 discussion
h1('7. Discussion')
h2('7.1 What the results show — and what they do not')
para('The results show that, on this simulated node, cross-checking related signals detects single-channel faults reliably, '
     'explains them in concrete terms, and keeps alert volume low once limits reflect real meter error. They do not show real-world '
     'accuracy: the faults were injected by us, they change one channel while the related channels stay correct, and the meter model, '
     'although realistic in magnitude, is itself an assumption.')
h2('7.2 Genuine events versus faulty data')
para('A disagreement between related channels supports a data-quality hypothesis; a coherent change across several channels without '
     'a physical mismatch is more likely a genuine operating event and is labelled as such. Telling them apart reliably needs trusted '
     'context (operator logs, weather, dispatch schedules) and independently labelled cases, which Phase 1 does not have. A '
     'coordinated attack that alters all related channels consistently would also pass the physics checks; this is an open problem.')
h2('7.3 Limitations')
for t in ['Building load, context, dispatch and electrical channels are synthetic or derived; PV/wind provenance and units are inherited assumptions.',
          'Meter error is modelled as independent Gaussian noise; real meters also have offsets, temperature drift and correlated errors.',
          'Faults are injected and single-channel; no confirmed real fault, cyberattack or genuine-event label is available.',
          'Hourly data only: sub-hourly transients and power-quality events are out of reach.',
          'The evidence score is a transparent heuristic, not a calibrated probability; explanation quality has not been rated by operators.',
          'The spike rule needs the next sample, so spikes are confirmed with a one-hour delay.']:
    bullet(t)

# ------------------------------------------------------------------ 8 goals
h1('8. Status against the four challenge goals')
table(['Challenge goal', 'Phase 1 delivery', 'Status'], [
    ['1. Anomaly & Bias Detection', 'Rules + six physics cross-checks with learned limits; six fault classes at three severities; '
     f'F1 {sel.f1:.3f}, {n_detected}/{n_events} events.', 'Implemented, simulated validation'],
    ['2. Confidence Scores & Explainability', 'Per-alert checks, mismatch vs limit, interpretation, suggested check, evidence score.',
     'Implemented; score not calibrated'],
    ['3. Human-AI Escalation & Governance', 'Review queue, four decision types, escalation policy, append-only audit log, no automatic action.',
     'Prototype'],
    ['4. Shared Learning Across the Noosphere', 'Design: shared pattern summaries, local-only data, compatibility checks, provenance, '
     'shadow testing; poisoning and privacy risks.', 'Design only (allowed in Phase 1)'],
], [1.9, 3.3, 1.4], 'Coverage of the official challenge goals.')

# ------------------------------------------------------------------ 9 conclusion
h1('9. Conclusion and next steps')
para(f'WattWitness shows that a small set of transparent checks — physical consistency with limits learned from data, plus '
     f'simple data-quality rules — can flag untrustworthy microgrid readings, explain why, and leave the final judgement to a person. '
     f'On realistically imperfect simulated telemetry it detected all {n_events} injected events with few false alerts, and the '
     f'sensitivity study makes its limits explicit. Next steps, in order:')
for t in ['Confirm units, battery-state meaning and sign conventions against the original simulation model.',
          'Replace the modelled meter error with equipment specifications or measured noise, including offsets and correlated error.',
          'Add labelled genuine operating events and evaluate the event-versus-fault interpretation; calibrate the evidence score.',
          'Run a replay test on recorded data from an independent node, without any connection to control equipment.',
          'Prototype the Goal 4 exchange between two independently configured nodes in shadow mode.']:
    bullet(t)

# ------------------------------------------------------------------ references + appendices
h1('References')
for i, ref in enumerate([
    'A. Abur and A. Gómez Expósito, Power System State Estimation: Theory and Implementation. New York: Marcel Dekker, 2004.',
    'F. T. Liu, K. M. Ting and Z.-H. Zhou, "Isolation Forest," in Proc. 8th IEEE Int. Conf. on Data Mining (ICDM), 2008, pp. 413–422.',
    'F. Pedregosa et al., "Scikit-learn: Machine Learning in Python," Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.',
    'IEC 62053-22, Electricity metering equipment — Particular requirements — Static meters for AC active energy (classes 0.1 S, 0.2 S and 0.5 S).',
    'National Institute of Standards and Technology, Artificial Intelligence Risk Management Framework (AI RMF 1.0), NIST AI 100-1, 2023.',
    'T. Brown, J. Hörsch and D. Schlachtberger, "PyPSA: Python for Power System Analysis," Journal of Open Research Software, vol. 6, no. 1, 2018.'], 1):
    p = doc.add_paragraph(f'[{i}] {ref}'); p.paragraph_format.left_indent = Inches(.3); p.paragraph_format.first_line_indent = Inches(-.3)

h1('Appendix A. Event-level results')
table(['Event', 'Type', 'Severity', 'Altered', 'Flagged', 'First flag after'],
      [[r.event_id, r.fault_type, r.severity, r.altered_rows, r.flagged_rows,
        '—' if pd.isna(r.hours_to_first_flag) else f'{int(r.hours_to_first_flag)} h'] for r in events.itertuples()],
      [.55, 1.1, 2.75, .65, .65, .9], 'Each injected event and how it was detected.', size=8)

h1('Appendix B. Lessons from the first iteration')
para('Reviewing the first version of the prototype exposed several weaknesses, all corrected in this version:')
for t in ['Physics checks were validated on noise-free derived signals, which made them trivially perfect → meter model and learned limits (Sections 3.4, 4.3).',
          'The spike rule fired on ordinary battery and grid dispatch steps → spike-and-return rule (Section 4.2).',
          'Hours whose value an injection did not change were labelled as faults → only altered readings are labelled; event-level metrics added.',
          'All injected spikes and biases were gross → severities now range from gross to subtle (Table 5).',
          'The missing-value rule covered 8 of 17 channels → all channels.',
          'A decision-log read error could lead to the log being overwritten, and a decision in one scenario marked the same hour reviewed in the other → append-only log keyed by scenario.']:
    bullet(t)

h1('Appendix C. Reproducibility')
para('All steps are scripted and seeded (meter model 7203, fault schedule 4152026, Isolation Forest 20261015). From the project '
     'folder: install requirements.txt, then run python scripts/run_phase1_pipeline.py --with-reports to regenerate the measured '
     'telemetry, injected faults, detector outputs, sensitivity study, figures and this report; python -m unittest discover -s tests '
     'runs the automated checks. The ideal dataset can be rebuilt from a permitted local export with '
     'scripts/build_simulated_microgrid_dataset.py. Data dictionary, assumptions, limitations and the cross-node design are in docs/.')

doc.core_properties.title = 'WattWitness - Phase 1 Technical Report'
doc.core_properties.subject = 'Explainable detection of untrustworthy microgrid telemetry on a simulated node'
doc.core_properties.author = 'WattWitness team'
doc.save(REPORT)
print(REPORT)
